import io
import logging
import zipfile
from pathlib import Path

from PIL import Image
from pdf2image import convert_from_bytes
from pypdf import PdfReader, PdfWriter

from ..errors import UserInputError
from ..utils.file_utils import save_upload, temp_job
from ..utils.pdf_render import iter_page_images

logger = logging.getLogger(__name__)


def pdf_to_images(file_bytes: bytes, filename: str, fmt: str = "png", dpi: int = 200) -> Path:
    with temp_job("convert") as job_dir:
        pdf_path = save_upload(file_bytes, filename, job_dir)

        output_dir = job_dir / "output"
        output_dir.mkdir(exist_ok=True)

        ext = fmt.lower()
        fmt_map = {"png": "PNG", "jpg": "JPEG", "jpeg": "JPEG", "tiff": "TIFF", "bmp": "BMP"}
        pillow_fmt = fmt_map.get(ext, "PNG")
        if ext not in fmt_map:
            ext = "png"

        saved = []
        for i, img in enumerate(iter_page_images(pdf_path, dpi=dpi)):
            out_path = output_dir / f"page_{i + 1}.{ext}"
            img.save(out_path, pillow_fmt)
            saved.append(out_path)

        zip_path = job_dir / "output.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for p in saved:
                zf.write(p, p.name)

        return zip_path


def images_to_pdf(files: list[tuple[bytes, str]], job_dir: Path | None = None) -> Path:
    if job_dir is None:
        with temp_job("convert") as job_dir:
            return _images_to_pdf(files, job_dir)
    return _images_to_pdf(files, job_dir)


def _images_to_pdf(files: list[tuple[bytes, str]], job_dir: Path) -> Path:
    writer = PdfWriter()

    for file_bytes, filename in files:
        img = Image.open(io.BytesIO(file_bytes))
        if img.mode == "RGBA":
            img = img.convert("RGB")

        buf = io.BytesIO()
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.save(buf, format="PDF")
        buf.seek(0)

        reader = PdfReader(buf)
        for page in reader.pages:
            writer.add_page(page)

    output_path = job_dir / "converted.pdf"
    with open(output_path, "wb") as f:
        writer.write(f)

    return output_path


def pdf_to_word(file_bytes: bytes, filename: str) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt

    with temp_job("convert") as job_dir:
        pdf_path = save_upload(file_bytes, filename, job_dir)

        reader = PdfReader(str(pdf_path))
        doc = Document()

        style = doc.styles["Normal"]
        style.font.size = Pt(11)

        has_text = False
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()

            if text:
                has_text = True
                for line in text.split("\n"):
                    line = line.strip()
                    if line:
                        doc.add_paragraph(line)
            else:
                try:
                    page_bytes = io.BytesIO()
                    single_writer = PdfWriter()
                    single_writer.add_page(page)
                    single_writer.write(page_bytes)
                    page_bytes.seek(0)
                    images = convert_from_bytes(page_bytes.read(), dpi=200)
                except Exception:
                    logger.warning("Could not rasterise page %d for DOCX output", i + 1, exc_info=True)
                    images = []

                if images:
                    img_path = job_dir / f"page_{i + 1}.png"
                    images[0].save(img_path, "PNG")
                    doc.add_picture(str(img_path), width=Inches(6))

            if i < len(reader.pages) - 1:
                doc.add_page_break()

        if not has_text:
            doc = Document()
            for i, page in enumerate(reader.pages):
                try:
                    page_bytes = io.BytesIO()
                    single_writer = PdfWriter()
                    single_writer.add_page(page)
                    single_writer.write(page_bytes)
                    page_bytes.seek(0)
                    images = convert_from_bytes(page_bytes.read(), dpi=200)
                    if images:
                        img_path = job_dir / f"page_{i + 1}.png"
                        images[0].save(img_path, "PNG")
                        doc.add_picture(str(img_path), width=Inches(6))
                        if i < len(reader.pages) - 1:
                            doc.add_page_break()
                except Exception:
                    pass

        output_path = job_dir / "converted.docx"
        doc.save(str(output_path))
        return output_path


def word_to_pdf(file_bytes: bytes, filename: str) -> Path:
    from docx import Document

    with temp_job("convert") as job_dir:
        docx_path = save_upload(file_bytes, filename, job_dir)

        try:
            doc = Document(str(docx_path))
        except Exception:
            raise UserInputError(
                "The Word document could not be read. It may be damaged or not a .docx file."
            ) from None

        from fpdf import FPDF
        from fpdf.enums import XPos, YPos
        from fpdf.errors import FPDFUnicodeEncodingException
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)

        tmp_dir = job_dir / "tmp_img"
        tmp_dir.mkdir(exist_ok=True)

        images = []
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                img_data = rel.target_part.blob
                ext = Path(rel.target_part.partname).suffix
                img_path = tmp_dir / f"img_{len(images)}{ext}"
                img_path.write_bytes(img_data)
                images.append(img_path)

        if images:
            return _images_to_pdf(
                [(img_path.read_bytes(), img_path.name) for img_path in images],
                job_dir,
            )

        try:
            for para in doc.paragraphs:
                if para.text.strip():
                    if pdf.pages_count == 0 or pdf.get_y() > 270:
                        pdf.add_page()
                    pdf.set_font("Helvetica", size=11)
                    # Return to the left margin on the next line; fpdf2's default leaves the
                    # cursor at the right edge, so a second paragraph had zero width.
                    pdf.multi_cell(0, 10, para.text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        except FPDFUnicodeEncodingException:
            raise UserInputError(
                "This document contains characters (for example non-Latin scripts or "
                "typographic quotes) that Word to PDF conversion does not support yet.",
                status_code=422,
            ) from None

        if pdf.pages_count == 0:
            pdf.add_page()

        output_path = job_dir / "converted.pdf"
        pdf.output(str(output_path))
        return output_path
